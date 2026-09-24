from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import re
import sys
from importlib.metadata import version
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation


TASK_ROOT = Path(__file__).resolve().parent
REPO_ROOT = TASK_ROOT.parents[1]
OUTPUT_ROOT = TASK_ROOT / "输出"
TEXT_ROOT = OUTPUT_ROOT / "完整文本_30条"
PACKET_ROOT = REPO_ROOT / "Q1" / "03_结果" / "Q1.1" / "v1" / "review_packet"
Q11_SCORES = REPO_ROOT / "Q1" / "03_结果" / "Q1.1" / "v1" / "sample_scores.csv.gz"
Q12_SCORES = REPO_ROOT / "Q1" / "03_结果" / "Q1.2" / "v1" / "sample_scores.csv.gz"
MAPPING = PACKET_ROOT / "restricted_id_mapping.csv"
REVIEW_MD = PACKET_ROOT / "blind_review_samples.md"
SCORE_INPUT = TASK_ROOT / "AI代理评分录入.csv"
WORKBOOK = OUTPUT_ROOT / "Q1.1_30条_AI辅助代理评分.xlsx"
PRIVATE_MANIFEST = OUTPUT_ROOT / "Q1.1_30条抽样清单_私有.csv"
PUBLIC_INDEX = OUTPUT_ROOT / "Q1.1_30条盲评索引.csv"
SUMMARY = OUTPUT_ROOT / "Q1.1_AI辅助代理评分说明.md"
REPRO = OUTPUT_ROOT / "复现清单.json"

SAMPLE_SEED = 20260925
ORDER_SEED = 20260926
DIMENSIONS = ["教育用途", "可读性", "连贯性", "信息传递", "整体质量"]
DIMENSION_KEYS = ["education", "readability", "coherence", "information", "overall"]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def stable_tie(blind_id: str) -> str:
    return hashlib.sha256(f"{SAMPLE_SEED}:{blind_id}".encode("utf-8")).hexdigest()


def parse_review_texts() -> dict[str, str]:
    raw = REVIEW_MD.read_text(encoding="utf-8")
    parts = re.split(r"(?m)^## (Q11-[0-9]+)\s*$", raw)
    texts: dict[str, str] = {}
    for index in range(1, len(parts), 2):
        blind_id, block = parts[index], parts[index + 1]
        lines = block.lstrip("\r\n").splitlines()
        if not lines or not lines[0].endswith("text"):
            raise ValueError(f"{blind_id} 缺少文本围栏")
        fence = lines[0][:-4]
        try:
            end = lines.index(fence, 1)
        except ValueError as exc:
            raise ValueError(f"{blind_id} 文本围栏未闭合") from exc
        texts[blind_id] = "\n".join(lines[1:end])
    if len(texts) != 793:
        raise ValueError(f"盲评原文数量异常：{len(texts)} != 793")
    return texts


def load_frame() -> pd.DataFrame:
    mapping = pd.read_csv(MAPPING)
    q11 = pd.read_csv(
        Q11_SCORES,
        usecols=["dataset", "source_domain", "id", "q_equal", "q_huber"],
    )
    q12 = pd.read_csv(
        Q12_SCORES,
        usecols=["dataset", "source_domain", "id", "conflict_intensity", "high_conflict"],
    )
    q11 = q11[q11["dataset"].str.lower().eq("a1")].copy()
    q12 = q12[q12["dataset"].str.lower().eq("a1")].copy()
    frame = mapping.merge(
        q11,
        left_on=["source_domain", "raw_id"],
        right_on=["source_domain", "id"],
        validate="one_to_one",
    ).merge(
        q12[["source_domain", "id", "conflict_intensity", "high_conflict"]],
        on=["source_domain", "id"],
        validate="one_to_one",
    )
    if len(frame) != 793 or frame["blind_id"].nunique() != 793:
        raise ValueError("793 条抽样框未能一一联接")
    frame["q_equal_percentile"] = frame["q_equal"].rank(method="average", pct=True)
    frame["q_huber_percentile"] = frame["q_huber"].rank(method="average", pct=True)
    frame["candidate_rank_gap"] = (
        frame["q_equal_percentile"] - frame["q_huber_percentile"]
    ).abs()
    frame["candidate_mean_percentile"] = (
        frame["q_equal_percentile"] + frame["q_huber_percentile"]
    ) / 2
    frame["stable_tie"] = frame["blind_id"].map(stable_tie)
    return frame


def representative_ids(frame: pd.DataFrame) -> list[str]:
    cells: list[dict[str, object]] = []
    for (domain, tertile), group in frame.groupby(
        ["source_domain", "q_equal_tertile"], sort=True
    ):
        cells.append({"domain": domain, "tertile": int(tertile), "n": len(group)})
    if len(cells) > 24:
        raise ValueError("非空单元超过 representative 名额")
    remaining_slots = 24 - len(cells)
    total = sum(int(cell["n"]) for cell in cells)
    for cell in cells:
        quota = int(cell["n"]) / total * remaining_slots
        cell["quota"] = quota
        cell["extra"] = int(quota)
    seats_left = remaining_slots - sum(int(cell["extra"]) for cell in cells)
    order = sorted(
        cells,
        key=lambda cell: (
            -(float(cell["quota"]) - int(cell["extra"])),
            str(cell["domain"]),
            int(cell["tertile"]),
        ),
    )
    for cell in order[:seats_left]:
        cell["extra"] = int(cell["extra"]) + 1
    rng = random.Random(SAMPLE_SEED)
    selected: list[str] = []
    for cell in sorted(cells, key=lambda item: (str(item["domain"]), int(item["tertile"]))):
        ids = sorted(
            frame.loc[
                frame["source_domain"].eq(cell["domain"])
                & frame["q_equal_tertile"].eq(cell["tertile"]),
                "blind_id",
            ].tolist()
        )
        rng.shuffle(ids)
        selected.extend(ids[: 1 + int(cell["extra"])])
    if len(selected) != 24 or len(set(selected)) != 24:
        raise ValueError("representative 抽样未得到 24 个唯一样本")
    return selected


def diagnostic_ids(frame: pd.DataFrame, excluded: set[str]) -> tuple[list[str], dict[str, str]]:
    pool = frame[~frame["blind_id"].isin(excluded)].copy()
    selected: list[str] = []
    labels: dict[str, str] = {}

    def take(category: str, candidates: pd.DataFrame, n: int, columns: list[str], ascending: list[bool]) -> None:
        available = candidates[~candidates["blind_id"].isin(selected)]
        rows = available.sort_values(columns, ascending=ascending, kind="mergesort").head(n)
        for blind_id in rows["blind_id"].tolist():
            selected.append(blind_id)
            labels[blind_id] = category

    take("candidate_disagreement", pool, 2, ["candidate_rank_gap", "stable_tie"], [False, True])
    take(
        "q1_2_high_conflict",
        pool[pool["high_conflict"].astype(bool)],
        2,
        ["conflict_intensity", "stable_tie"],
        [False, True],
    )
    take("extreme_high_q", pool, 1, ["candidate_mean_percentile", "stable_tie"], [False, True])
    take("extreme_low_q", pool, 1, ["candidate_mean_percentile", "stable_tie"], [True, True])
    if len(selected) < 6:
        take("fallback", pool, 6 - len(selected), ["stable_tie"], [True])
    if len(selected) != 6 or len(set(selected)) != 6:
        raise ValueError("diagnostic 抽样未得到 6 个唯一样本")
    return selected, labels


def selected_frame() -> tuple[pd.DataFrame, dict[str, str]]:
    frame = load_frame()
    representative = representative_ids(frame)
    diagnostic, labels = diagnostic_ids(frame, set(representative))
    ids = representative + diagnostic
    selected = frame.set_index("blind_id").loc[ids].reset_index()
    selected["sample_type"] = ["representative"] * 24 + ["diagnostic"] * 6
    selected["diagnostic_category"] = selected["blind_id"].map(labels).fillna("")
    selected["selection_order"] = range(1, 31)
    return selected, labels


def write_workbook(index: pd.DataFrame, scores: pd.DataFrame | None = None) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "评分"
    headers = ["评审顺序", "blind_id", "完整文本文件", "字符数", *DIMENSIONS, "评分来源", "备注"]
    ws.append(headers)
    score_lookup = {} if scores is None else scores.set_index("blind_id").to_dict("index")
    for row in index.itertuples(index=False):
        values = score_lookup.get(row.blind_id, {})
        ws.append(
            [
                int(row.review_order),
                row.blind_id,
                row.text_file,
                int(row.char_count),
                *[values.get(key, "") for key in DIMENSION_KEYS],
                "AI辅助代理评分（非人工、非独立盲评）" if values else "待评分",
                values.get("notes", ""),
            ]
        )
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{ws.max_row}"
    widths = [10, 14, 38, 12, 11, 11, 11, 11, 11, 34, 42]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + col)].width = width
    validation = DataValidation(type="whole", operator="between", formula1="1", formula2="5")
    validation.error = "只能填写 1–5 的整数"
    validation.errorTitle = "评分无效"
    ws.add_data_validation(validation)
    validation.add(f"E2:I{ws.max_row}")
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    scale = wb.create_sheet("评分量表")
    scale.append(["项目", "说明"])
    scale_rows = [
        ("评分范围", "1=很弱/基本不可用；2=偏弱；3=中等/可用；4=较强；5=很强。"),
        ("教育用途", "文本用于学习、教学、知识获得或技能训练的价值。"),
        ("可读性", "语言、结构和格式是否便于目标读者理解。"),
        ("连贯性", "内容是否围绕明确主题组织，前后关系是否自然。"),
        ("信息传递", "信息是否清楚、充分且有效地传达给读者。"),
        ("整体质量", "综合判断；这是主评审效标。"),
        ("来源声明", "本文件为用户委托的 AI 辅助代理评分，不是人类评审，也不是第二名独立评审。"),
        ("盲评限制", "执行者在评分前已接触抽样调试元数据，因此本结果不得称为严格盲评。"),
        ("长文本处理", "Excel 单元格有 32,767 字符上限；完整正文保存在旁侧 TXT 文件，并由 SHA-256 绑定。"),
    ]
    for item in scale_rows:
        scale.append(item)
    for cell in scale[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    scale.column_dimensions["A"].width = 18
    scale.column_dimensions["B"].width = 105
    for row in scale.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    integrity = wb.create_sheet("完整性索引")
    integrity.append(["blind_id", "完整文本文件", "字符数", "review_text_sha256"])
    for row in index.itertuples(index=False):
        integrity.append([row.blind_id, row.text_file, int(row.char_count), row.review_text_sha256])
    for cell in integrity[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    integrity.freeze_panes = "A2"
    integrity.column_dimensions["A"].width = 14
    integrity.column_dimensions["B"].width = 40
    integrity.column_dimensions["C"].width = 12
    integrity.column_dimensions["D"].width = 70
    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK)


def prepare() -> None:
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    TEXT_ROOT.mkdir(parents=True, exist_ok=True)
    texts = parse_review_texts()
    selected, _ = selected_frame()
    mapping_ids = set(pd.read_csv(MAPPING, usecols=["blind_id"])["blind_id"])
    if set(texts) != mapping_ids:
        missing_text = sorted(mapping_ids - set(texts))
        extra_text = sorted(set(texts) - mapping_ids)
        raise ValueError(
            "793 条盲评文本与 restricted mapping 的 blind_id 集合不一致："
            f"缺失={missing_text[:5]}，多余={extra_text[:5]}"
        )
    rows: list[dict[str, object]] = []
    for row in selected.itertuples(index=False):
        text = texts[row.blind_id]
        text_path = TEXT_ROOT / f"{row.blind_id}.txt"
        text_path.write_text(text, encoding="utf-8", newline="\n")
        rows.append(
            {
                "blind_id": row.blind_id,
                "sample_type": row.sample_type,
                "diagnostic_category": row.diagnostic_category,
                "source_domain": row.source_domain,
                "raw_id": row.raw_id,
                "q_equal_tertile": int(row.q_equal_tertile),
                "q_equal": row.q_equal,
                "q_huber": row.q_huber,
                "conflict_intensity": row.conflict_intensity,
                "high_conflict": bool(row.high_conflict),
                "source_content_sha256": row.content_sha256,
                "review_text_sha256": sha256_file(text_path),
                "char_count": len(text),
                "text_file": str(text_path.relative_to(OUTPUT_ROOT)).replace("\\", "/"),
                "selection_order": int(row.selection_order),
            }
        )
    manifest = pd.DataFrame(rows)
    manifest.to_csv(PRIVATE_MANIFEST, index=False, encoding="utf-8-sig")
    review_order = manifest["blind_id"].tolist()
    random.Random(ORDER_SEED).shuffle(review_order)
    order_map = {blind_id: index + 1 for index, blind_id in enumerate(review_order)}
    public = manifest[["blind_id", "review_text_sha256", "char_count", "text_file"]].copy()
    public["review_order"] = public["blind_id"].map(order_map)
    public = public.sort_values("review_order")
    public.to_csv(PUBLIC_INDEX, index=False, encoding="utf-8-sig")
    write_workbook(public)
    hashes = {
        str(path.relative_to(REPO_ROOT)).replace("\\", "/"): sha256_file(path)
        for path in [MAPPING, REVIEW_MD, Q11_SCORES, Q12_SCORES]
    }
    script_rel = str(Path(__file__).resolve().relative_to(REPO_ROOT)).replace("\\", "/")
    python_executable = str(Path(sys.executable).resolve())
    command_prefix = f'"{python_executable}" -B "{script_rel}"'
    REPRO.write_text(
        json.dumps(
            {
                "schema_version": "q1-1-ai-proxy-review-v1",
                "rating_source": "AI-assisted proxy; not human; not independent blind review",
                "sample_seed": SAMPLE_SEED,
                "review_order_seed": ORDER_SEED,
                "input_sha256": hashes,
                "sample_count": 30,
                "representative_count": 24,
                "diagnostic_count": 6,
                "runtime": {
                    "python_executable": python_executable,
                    "python_version": sys.version,
                    "pandas": version("pandas"),
                    "openpyxl": version("openpyxl"),
                },
                "script_sha256": sha256_file(Path(__file__).resolve()),
                "working_directory": str(REPO_ROOT),
                "prepare_command": f"{command_prefix} prepare",
                "apply_command": f"{command_prefix} apply",
                "verify_command": f"{command_prefix} verify",
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    write_summary(manifest, scored=False)


def read_scores() -> pd.DataFrame:
    if not SCORE_INPUT.exists():
        raise FileNotFoundError(f"缺少评分录入文件：{SCORE_INPUT}")
    scores = pd.read_csv(SCORE_INPUT, dtype={"blind_id": str})
    required = ["blind_id", *DIMENSION_KEYS, "notes"]
    if scores.columns.tolist() != required:
        raise ValueError(f"评分列应为：{required}")
    if len(scores) != 30 or scores["blind_id"].nunique() != 30:
        raise ValueError("评分文件必须含 30 个唯一 blind_id")
    for key in DIMENSION_KEYS:
        numeric = pd.to_numeric(scores[key], errors="raise")
        if not numeric.between(1, 5).all() or not (numeric % 1 == 0).all():
            raise ValueError(f"{key} 必须全部为 1–5 整数")
        scores[key] = numeric.astype(int)
    return scores


def apply_scores() -> None:
    if not PUBLIC_INDEX.exists():
        prepare()
    index = pd.read_csv(PUBLIC_INDEX)
    scores = read_scores()
    if set(index["blind_id"]) != set(scores["blind_id"]):
        raise ValueError("评分 blind_id 集合与抽样集合不一致")
    write_workbook(index, scores)
    manifest = pd.read_csv(PRIVATE_MANIFEST)
    write_summary(manifest, scored=True)
    reproduction = json.loads(REPRO.read_text(encoding="utf-8"))
    reproduction["score_input_sha256"] = sha256_file(SCORE_INPUT)
    reproduction["outputs_sha256"] = {
        str(path.relative_to(TASK_ROOT)).replace("\\", "/"): sha256_file(path)
        for path in [PRIVATE_MANIFEST, PUBLIC_INDEX, WORKBOOK, SUMMARY]
    }
    reproduction["text_content_set_sha256"] = hashlib.sha256(
        "\n".join(
            sorted(
                f"{path.name}:{sha256_file(path)}"
                for path in TEXT_ROOT.glob("Q11-*.txt")
            )
        ).encode("utf-8")
    ).hexdigest()
    REPRO.write_text(
        json.dumps(reproduction, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    verify(require_scores=True)


def write_summary(manifest: pd.DataFrame, scored: bool) -> None:
    over_limit = int((manifest["char_count"] > 32767).sum())
    score_section = ""
    if scored:
        scores = read_scores()
        means = {key: float(scores[key].mean()) for key in DIMENSION_KEYS}
        counts = scores["overall"].value_counts().reindex(range(1, 6), fill_value=0)
        score_section = f"""
## 描述性评分结果

- 教育用途均值：{means['education']:.2f}
- 可读性均值：{means['readability']:.2f}
- 连贯性均值：{means['coherence']:.2f}
- 信息传递均值：{means['information']:.2f}
- 整体质量均值：{means['overall']:.2f}
- 整体质量分布（1–5 分）：{', '.join(f'{score}分={int(counts[score])}条' for score in range(1, 6))}

以上仅为本次 30 条 AI 代理核读的描述统计，不计算双评一致性，也不进入 q_equal/q_huber 的人工效度比较。
"""
    text = f"""# Q1.1 30 条 AI 辅助代理评分说明

## 状态

- 抽样：24 条 representative + 6 条 diagnostic，共 30 条。
- 评分状态：{'已完成 30×5 项代理评分。' if scored else '等待代理评分。'}
- 评分来源：用户委托的 AI 辅助代理评分；不是人类评审，也不是第二名独立评审。
- 盲评限制：执行者在评分前接触过抽样调试元数据，因此本结果不能称为严格盲评。

## 文本承载

30 条文本合计 {int(manifest['char_count'].sum()):,} 个字符；其中 {over_limit} 条超过 Excel 单元格 32,767 字符上限。为避免截断，完整正文分别保存在 `完整文本_30条/*.txt`，工作簿通过文件名、字符数和 SHA-256 与正文绑定。

## 评分执行方式

短文本直接全文核读；长文本采用首部、25%、50%、75%位置与尾部的等距分点核读，并结合字符数、行数和重复行比例判断格式及信息质量。评分备注逐条记录了全文核读或分点核读方式。
{score_section}

## 使用边界

本结果只能作为探索性 AI 代理核读记录，不能替代附件要求的两名独立人工评审，不能据此声称“Q1.1 已完成人工效度验证”或直接选择 q_equal/q_huber 的胜者。评分者未查看工作簿中的候选分数，但因调试阶段见过部分来源域和诊断标签，严格盲法已被破坏。

## 复现命令

以下命令从仓库根目录执行；实际解释器和完整命令同时记录在 `复现清单.json`。

```powershell
& '<workspace-python>' -B 'zwj/Q1.1_30条小样本代理评审_20260924/q1_1_proxy_review.py' prepare
& '<workspace-python>' -B 'zwj/Q1.1_30条小样本代理评审_20260924/q1_1_proxy_review.py' apply
& '<workspace-python>' -B 'zwj/Q1.1_30条小样本代理评审_20260924/q1_1_proxy_review.py' verify
```
"""
    SUMMARY.write_text(text, encoding="utf-8")


def verify(require_scores: bool = False) -> None:
    reproduction = json.loads(REPRO.read_text(encoding="utf-8"))
    current_script_hash = sha256_file(Path(__file__).resolve())
    if reproduction.get("script_sha256") != current_script_hash:
        raise ValueError(
            "当前脚本与复现清单未绑定："
            f"manifest={reproduction.get('script_sha256')} current={current_script_hash}"
        )
    manifest = pd.read_csv(PRIVATE_MANIFEST)
    public = pd.read_csv(PUBLIC_INDEX)
    if len(manifest) != 30 or manifest["blind_id"].nunique() != 30:
        raise ValueError("私有清单不是 30 个唯一样本")
    if (manifest["sample_type"] == "representative").sum() != 24:
        raise ValueError("representative 数量不是 24")
    if (manifest["sample_type"] == "diagnostic").sum() != 6:
        raise ValueError("diagnostic 数量不是 6")
    if set(manifest["blind_id"]) != set(public["blind_id"]):
        raise ValueError("公开索引与私有清单的 blind_id 集合不同")
    for row in public.itertuples(index=False):
        path = OUTPUT_ROOT / row.text_file
        if not path.exists() or sha256_file(path) != row.review_text_sha256:
            raise ValueError(f"正文缺失或哈希不符：{row.blind_id}")
        if len(path.read_text(encoding="utf-8")) != int(row.char_count):
            raise ValueError(f"正文字数不符：{row.blind_id}")
    wb = load_workbook(WORKBOOK, data_only=False)
    if wb.sheetnames != ["评分", "评分量表", "完整性索引"]:
        raise ValueError("工作簿工作表结构异常")
    ws = wb["评分"]
    if ws.max_row != 31:
        raise ValueError("工作簿评分行数不是 30")
    if require_scores:
        if reproduction.get("score_input_sha256") != sha256_file(SCORE_INPUT):
            raise ValueError("评分录入文件与复现清单哈希不一致")
        for relative, expected in reproduction.get("outputs_sha256", {}).items():
            path = TASK_ROOT / relative
            if not path.exists() or sha256_file(path) != expected:
                raise ValueError(f"输出文件与复现清单哈希不一致：{relative}")
        for row in ws.iter_rows(min_row=2, min_col=5, max_col=9, values_only=True):
            if any(value not in {1, 2, 3, 4, 5} for value in row):
                raise ValueError("工作簿存在空评分或越界评分")
    print(
        json.dumps(
            {
                "status": "PASS",
                "sample_count": 30,
                "representative_count": 24,
                "diagnostic_count": 6,
                "text_sha256_pass": 30,
                "scores_complete": require_scores,
                "workbook": str(WORKBOOK),
            },
            ensure_ascii=False,
        )
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("command", choices=["prepare", "apply", "verify"])
    args = parser.parse_args()
    if args.command == "prepare":
        prepare()
        verify(require_scores=False)
    elif args.command == "apply":
        apply_scores()
    else:
        verify(require_scores=SCORE_INPUT.exists())


if __name__ == "__main__":
    main()
