"""D 题（山区洪涝灾害下无人机运输与通信协同优化）真实参数加载层。

数据来源（全部为附件原值，不做任何假设性修改）：

- ``中文题目/D题/数据/无人机应急物资运输基础数据/*.xlsx``
- ``中文题目/D题/数据/镇龙乡地理空间数据/.../镇龙乡及周边30米DEM.mat``

附录未明写的计算口径集中记录在 :mod:`physics` 与 :mod:`comm` 的模块注释中，
本模块只负责"把附件数值读进来"。
"""

from __future__ import annotations

import math
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
TOPIC = ROOT / "中文题目" / "D题"
DATA = TOPIC / "数据"
BASE_DATA = DATA / "无人机应急物资运输基础数据"
GEO_DATA = DATA / "镇龙乡地理空间数据"
DEM_MAT = (
    GEO_DATA
    / "镇龙乡及周边地理数据"
    / "数字高程模型数据（DEM）"
    / "镇龙乡及周边30米DEM.mat"
)

# ---------------------------------------------------------------- 基础结构


@dataclass(frozen=True)
class TransportModel:
    """运输无人机机型参数（“运输无人机数据.xlsx / 三类机型参数”）。"""

    id: str
    name: str
    empty_mass: float          # 含电池空载总质量 kg
    max_payload: float         # 最大载货质量 kg
    volume: float              # 可用装载体积 m^3
    cruise_speed: float        # 计划巡航速度 m/s
    range_empty: float         # 空载标准航程 m
    range_full: float          # 满载标准航程 m
    battery_energy: float      # 电池可用能量 kWh
    reserve_ratio: float       # 返航电量下限（比例）
    ground_prep: float         # 工位固定准备时间 s
    load_per_box: float        # 每箱装载时间 s
    handover_base: float       # 接收点基础交接时间 s
    handover_per_box: float    # 每箱增加交接时间 s
    v_up: float                # 最大爬升速度 m/s
    v_down: float              # 最大下降速度 m/s
    eta_up: float              # 爬升能耗效率
    eta_down: float            # 下降能耗效率（0 表示不单独计算）


@dataclass(frozen=True)
class RelayModel:
    """中继无人机机型参数（“中继无人机数据.xlsx / 中继机型参数”）。"""

    id: str
    name: str
    empty_mass: float          # 含能源组件空载总质量 kg
    comm_module_mass: float    # 中继通信模块质量 kg
    takeoff_mass: float        # 计划起飞总质量 kg
    cruise_speed: float        # 计划巡航速度 m/s
    cruise_power: float        # 巡航功率 kW
    energy: float              # 能源组件可用能量 kWh
    reserve_ratio: float       # 返航电量下限（比例）
    ground_prep: float         # 工位固定准备时间 s
    link_setup: float          # 建链时间 s
    turnaround: float          # 架次周转时间 s
    v_up: float                # 最大爬升速度 m/s
    v_down: float              # 最大下降速度 m/s
    eta_up: float              # 爬升能耗效率
    eta_down: float            # 下降能耗效率
    hover_power: float         # 悬停功率 kW
    comm_power: float          # 通信附加功率 kW
    max_hover_agl: float       # 最大悬停离地高度 m


@dataclass(frozen=True)
class Node:
    """调度中心或服务区节点。"""

    id: str
    name: str
    lon: float
    lat: float
    elevation: float           # 附件给出的地面海拔 m
    population: int = 0
    is_depot: bool = False


@dataclass(frozen=True)
class Box:
    """不可拆货箱。"""

    id: str
    area: str                  # 所属服务区
    kind: str                  # 医疗物资 / 饮用水 / 应急食品 / 生活卫生用品
    mass: float                # kg
    volume: float              # m^3
    is_first_batch: bool       # 是否首批保障
    first_deadline: float | None   # 首批截止时间 s
    expected_time: float       # 期望送达时间 s
    priority: float            # 应急优先系数


@dataclass(frozen=True)
class LinkParams:
    """通信链路参数（“通信链路参数.xlsx”）。"""

    freq_mhz: float
    l_sys: float
    l_obs: float
    p_sens: float              # dBm
    margin: float              # dB
    tx_transport: tuple[float, float]   # (Pt dBm, G dBi)
    tx_relay_access: tuple[float, float]
    tx_relay_backhaul: tuple[float, float]
    tx_gateway: tuple[float, float]
    gateway_agl: float         # 固定网关天线离地高度 m

    @property
    def p_threshold(self) -> float:
        """有效接收门限 P_th = P_sens + M。"""
        return self.p_sens + self.margin


@dataclass
class Scenario:
    """一个完整的 D 题算例（可为真实全量或压缩小实例）。"""

    depot: Node
    areas: dict[str, Node]
    boxes: list[Box]
    transport_models: dict[str, TransportModel]
    relay_model: RelayModel
    transport_uavs: dict[str, str]        # 无人机编号 -> 机型
    relay_uavs: list[str]
    battery_stock: dict[str, int]          # 机型 -> 组数
    battery_charge_full: dict[str, float]  # 机型 -> 等效完全充电时间 s
    relay_energy_stock: int
    relay_charge_full: float
    link: LinkParams
    dem: object = field(default=None, repr=False)

    # -- 便捷访问 -------------------------------------------------------
    def node(self, nid: str) -> Node:
        if nid == self.depot.id:
            return self.depot
        return self.areas[nid]

    def boxes_of(self, area: str) -> list[Box]:
        return [b for b in self.boxes if b.area == area]

    def uavs_of_model(self, mid: str) -> list[str]:
        return [u for u, m in self.transport_uavs.items() if m == mid]


# ---------------------------------------------------------------- xlsx 解析


def _clean_row(row: tuple) -> list:
    vals = list(row)
    while vals and (vals[-1] is None or str(vals[-1]).strip() == ""):
        vals.pop()
    return vals


def _is_header(row: tuple) -> bool:
    vals = [v for v in row if v is not None and str(v).strip() != ""]
    if len(vals) < 2:
        return False
    return all(isinstance(v, str) for v in vals)


def read_blocks(path: Path, sheet: str | None = None) -> dict[str, tuple[list[str], list[list]]]:
    """把“块标题 + 表头 + 数据行”结构的 sheet 解析为字典。

    无块标题的表（第一行即表头）以 ``""`` 作为键返回。
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    blocks: dict[str, tuple[list[str], list[list]]] = {}

    def consume(header_row: tuple, start: int) -> tuple[list[list], int]:
        data, j = [], start
        while j < len(rows):
            vals = _clean_row(rows[j])
            if not vals:
                break
            # 遇到下一个块标题（单值行 + 下一行是表头）则停止
            non_empty = [v for v in vals if v is not None and str(v).strip() != ""]
            if len(non_empty) == 1 and j + 1 < len(rows) and _is_header(rows[j + 1]):
                break
            data.append(vals)
            j += 1
        return data, j

    header_title = [str(v).strip() for v in _clean_row(rows[0])] if rows else []
    if _is_header(rows[0]) and len(header_title) >= 3:
        data, _ = consume(rows[0], 1)
        blocks[""] = (header_title, data)

    i = 0
    while i < len(rows):
        vals = _clean_row(rows[i])
        non_empty = [v for v in vals if v is not None and str(v).strip() != ""]
        if len(non_empty) == 1 and i + 1 < len(rows) and _is_header(rows[i + 1]):
            title = str(non_empty[0]).strip()
            header = [str(v).strip() for v in _clean_row(rows[i + 1])]
            data, j = consume(rows[i + 1], i + 2)
            blocks[title] = (header, data)
            i = max(j, i + 1)
            continue
        i += 1
    return blocks


def _col(header: list[str], *keys: str) -> int:
    for key in keys:
        for idx, h in enumerate(header):
            if key in h:
                return idx
    raise KeyError(f"未找到列 {keys}，实际表头：{header}")


def _rows_as_dicts(header: list[str], data: list[list]) -> list[dict]:
    out = []
    for r in data:
        out.append({header[i]: (r[i] if i < len(r) else None) for i in range(len(header))})
    return out


# ---------------------------------------------------------------- 加载函数


def load_transport_models() -> dict[str, TransportModel]:
    blocks = read_blocks(BASE_DATA / "运输无人机数据.xlsx")
    header, data = blocks["三类机型参数"]
    models: dict[str, TransportModel] = {}
    for r in data:
        mid = str(r[0]).strip()
        models[mid] = TransportModel(
            id=mid,
            name=str(r[1]).strip(),
            empty_mass=float(r[2]),
            max_payload=float(r[3]),
            volume=float(r[4]),
            cruise_speed=float(r[5]),
            range_empty=float(r[6]),
            range_full=float(r[7]),
            battery_energy=float(r[8]),
            reserve_ratio=float(r[9]) / 100.0,
            ground_prep=float(r[10]),
            load_per_box=float(r[11]),
            handover_base=float(r[12]),
            handover_per_box=float(r[13]),
            v_up=float(r[14]),
            v_down=float(r[15]),
            eta_up=float(r[16]),
            eta_down=float(r[17]),
        )
    return models


def load_transport_fleet() -> dict[str, str]:
    blocks = read_blocks(BASE_DATA / "运输无人机数据.xlsx")
    _, data = blocks["逐架无人机清单"]
    return {str(r[0]).strip(): str(r[1]).strip() for r in data}


def load_battery_stock() -> tuple[dict[str, int], dict[str, float]]:
    blocks = read_blocks(BASE_DATA / "运输无人机数据.xlsx")
    _, data = blocks["共享电池库存"]
    stock = {str(r[0]).strip(): int(r[1]) for r in data}
    full = {str(r[0]).strip(): float(r[2]) for r in data}
    return stock, full


def load_relay_model() -> RelayModel:
    blocks = read_blocks(BASE_DATA / "中继无人机数据.xlsx")
    _, data = blocks["中继机型参数"]
    r = data[0]
    return RelayModel(
        id=str(r[0]).strip(),
        name=str(r[1]).strip(),
        empty_mass=float(r[2]),
        comm_module_mass=float(r[3]),
        takeoff_mass=float(r[4]),
        cruise_speed=float(r[5]),
        cruise_power=float(r[6]),
        energy=float(r[7]),
        reserve_ratio=float(r[8]) / 100.0,
        ground_prep=float(r[9]),
        link_setup=float(r[10]),
        turnaround=float(r[11]),
        v_up=float(r[12]),
        v_down=float(r[13]),
        eta_up=float(r[14]),
        eta_down=float(r[15]),
        hover_power=float(r[16]),
        comm_power=float(r[17]),
        max_hover_agl=float(r[18]),
    )


def load_relay_fleet() -> list[str]:
    blocks = read_blocks(BASE_DATA / "中继无人机数据.xlsx")
    _, data = blocks["逐架中继无人机清单"]
    return [str(r[0]).strip() for r in data]


def load_relay_stock() -> tuple[int, float]:
    blocks = read_blocks(BASE_DATA / "中继无人机数据.xlsx")
    _, data = blocks["共享能源组件库存"]
    return int(data[0][1]), float(data[0][2])


def load_nodes() -> tuple[Node, dict[str, Node]]:
    blocks = read_blocks(BASE_DATA / "调度中心与服务区.xlsx")
    _, depot_rows = blocks["调度中心"]
    d = depot_rows[0]
    depot = Node(
        id=str(d[0]).strip(), name=str(d[1]).strip(), lon=float(d[2]),
        lat=float(d[3]), elevation=float(d[4]), is_depot=True,
    )
    areas: dict[str, Node] = {}
    _, area_rows = blocks["服务区"]
    for r in area_rows:
        nid = str(r[0]).strip()
        areas[nid] = Node(
            id=nid, name=str(r[1]).strip(), lon=float(r[2]), lat=float(r[3]),
            elevation=float(r[4]), population=int(r[5]) if r[5] is not None else 0,
        )
    return depot, areas


def load_boxes() -> list[Box]:
    blocks = read_blocks(BASE_DATA / "物资需求与配送时限.xlsx", sheet="逐箱货箱清单")
    header, data = blocks[""]
    idx = {name: i for i, name in enumerate(header)}
    boxes = []
    for r in data:
        def g(key: str):
            i = idx[key]
            return r[i] if i < len(r) else None

        boxes.append(Box(
            id=str(g("货箱编号")).strip(),
            area=str(g("服务区编号")).strip(),
            kind=str(g("物资类型")).strip(),
            mass=float(g("单箱质量（kg）")),
            volume=float(g("单箱体积（m³）")),
            is_first_batch=str(g("是否首批保障")).strip() == "是",
            first_deadline=(float(g("首批截止时间（s）")) if g("首批截止时间（s）") is not None else None),
            expected_time=float(g("期望送达时间（s）")),
            priority=float(g("应急优先系数")),
        ))
    return boxes


def load_link_params() -> LinkParams:
    path = BASE_DATA / "通信链路参数.xlsx"
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]

    def find(category: str, symbol: str) -> float:
        for r in rows:
            vals = [v for v in r if v is not None and str(v).strip() != ""]
            if len(vals) >= 3 and str(vals[0]).strip() == category and str(vals[1]).strip().startswith(symbol):
                return float(vals[-1])   # 参数值在最后（表格中间夹着“符号”列）
        raise KeyError(f"未找到参数 {category}/{symbol}")

    def pair(category: str) -> tuple[float, float]:
        pt = find(category, "发射功率")
        g = find(category, "天线增益")
        return pt, g

    return LinkParams(
        freq_mhz=find("传播参数", "载波频率"),
        l_sys=find("传播参数", "系统损耗"),
        l_obs=find("传播参数", "地形遮挡附加损耗"),
        p_sens=find("接收参数", "接收灵敏度"),
        margin=find("接收参数", "衰落裕量"),
        tx_transport=pair("运输无人机"),
        tx_relay_access=pair("中继接入端"),
        tx_relay_backhaul=pair("中继回传端"),
        tx_gateway=pair("固定网关 G01"),
        gateway_agl=find("固定网关 G01", "天线离地高度"),
    )


def load_scenario(dem=None) -> Scenario:
    """加载完整真实算例（15 服务区、80 货箱、8 架运输机、2 架中继机）。"""
    depot, areas = load_nodes()
    stock, full = load_battery_stock()
    relay_stock, relay_full = load_relay_stock()
    return Scenario(
        depot=depot,
        areas=areas,
        boxes=load_boxes(),
        transport_models=load_transport_models(),
        relay_model=load_relay_model(),
        transport_uavs=load_transport_fleet(),
        relay_uavs=load_relay_fleet(),
        battery_stock=stock,
        battery_charge_full=full,
        relay_energy_stock=relay_stock,
        relay_charge_full=relay_full,
        link=load_link_params(),
        dem=dem,
    )


# ------------------------------------------------- 方案 10.2 的统一小实例

#: 统一小实例使用的服务区（方案 10.2）
SMALL_AREAS = ["S001", "S002", "S015"]

#: 统一小实例使用的 8 个货箱，覆盖医疗物资/饮用水/应急食品/生活卫生用品四类
SMALL_BOXES = [
    "S001-MED-01", "S001-WAT-01", "S001-FOD-01", "S001-HYG-01",
    "S002-MED-01", "S002-WAT-01", "S002-HYG-01",
    "S015-WAT-01",
]

#: 统一小实例保留的运输无人机（1 架 A、1 架 B、1 架 C，方案 10.2）
SMALL_UAVS = ["U01", "U05", "U07"]

#: 统一小实例每机型保留的共享电池组数（方案 10.2 取 1~2 组，这里取 2 组）
SMALL_BATTERY_STOCK = {"A": 2, "B": 2, "C": 2}

#: 统一小实例保留的中继无人机与能源组件
SMALL_RELAY_UAVS = ["R01"]
SMALL_RELAY_ENERGY_STOCK = 2


def build_small_scenario(dem=None, *, areas=None, box_ids=None, uavs=None,
                         battery_stock=None, relay_uavs=None,
                         relay_energy_stock=None) -> Scenario:
    """按方案 10.2 压缩出的统一小实例（可覆写各子集）。"""
    full = load_scenario(dem=dem)
    area_ids = areas if areas is not None else SMALL_AREAS
    box_set = set(box_ids if box_ids is not None else SMALL_BOXES)
    uav_ids = uavs if uavs is not None else SMALL_UAVS
    stock = battery_stock if battery_stock is not None else dict(SMALL_BATTERY_STOCK)
    keep_models = {full.transport_uavs[u] for u in uav_ids}
    return Scenario(
        depot=full.depot,
        areas={a: full.areas[a] for a in area_ids},
        boxes=[b for b in full.boxes if b.id in box_set],
        transport_models={m: full.transport_models[m] for m in keep_models},
        relay_model=full.relay_model,
        transport_uavs={u: full.transport_uavs[u] for u in uav_ids},
        relay_uavs=list(relay_uavs if relay_uavs is not None else SMALL_RELAY_UAVS),
        battery_stock={m: stock[m] for m in keep_models},
        battery_charge_full={m: full.battery_charge_full[m] for m in keep_models},
        relay_energy_stock=relay_energy_stock if relay_energy_stock is not None else SMALL_RELAY_ENERGY_STOCK,
        relay_charge_full=full.relay_charge_full,
        link=full.link,
        dem=dem,
    )


# ---------------------------------------------------------------- 几何辅助


def haversine_m(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    """两点间水平大圆距离（m）。"""
    r = 6371008.8
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = p2 - p1
    dl = math.radians(lon2 - lon1)
    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * r * math.asin(math.sqrt(a))


if __name__ == "__main__":  # 手工核对参数用
    sc = load_scenario()
    print("机型:", {k: (v.max_payload, v.volume, v.battery_energy) for k, v in sc.transport_models.items()})
    print("机队:", sc.transport_uavs)
    print("电池:", sc.battery_stock, sc.battery_charge_full)
    print("中继:", sc.relay_uavs, sc.relay_energy_stock, sc.relay_charge_full)
    print("中继机型:", sc.relay_model)
    print("链路:", sc.link)
    print("节点数:", len(sc.areas), "货箱数:", len(sc.boxes))
    print("O01->S015 距离(m):", round(haversine_m(sc.depot.lat, sc.depot.lon,
                                                  sc.areas["S015"].lat, sc.areas["S015"].lon), 1))
