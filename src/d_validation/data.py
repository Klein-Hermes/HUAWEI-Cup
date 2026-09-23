"""Read the D-problem inputs and build the deliberately small validation case."""

from __future__ import annotations

from math import cos, radians
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .model import (
    Aircraft,
    AircraftType,
    Battery,
    Box,
    Node,
    Relay,
    RelayComponent,
    RelayType,
    Scenario,
    Terrain,
)


DATA_DIR = Path("中文题目/D题/数据/无人机应急物资运输基础数据")
GEO_DIR = Path("中文题目/D题/数据/镇龙乡地理空间数据/镇龙乡及周边地理数据")


def _rows(path: Path, sheet: str) -> list[list[Any]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[sheet]
    return [list(row) for row in worksheet.iter_rows(values_only=True)]


def _as_float(value: Any, default: float | None = None) -> float | None:
    if value is None or value == "":
        return default
    return float(value)


def _local_xy(lon: float, lat: float, lon0: float, lat0: float) -> tuple[float, float]:
    """Small-area equirectangular conversion; enough for the validator fixtures."""

    return (
        (lon - lon0) * 111_320.0 * cos(radians(lat0)),
        (lat - lat0) * 110_540.0,
    )


def _read_base_tables(root: Path) -> dict[str, Any]:
    data_root = root / DATA_DIR
    center_rows = _rows(data_root / "调度中心与服务区.xlsx", "数据")
    aircraft_rows = _rows(data_root / "运输无人机数据.xlsx", "数据")
    relay_rows = _rows(data_root / "中继无人机数据.xlsx", "数据")
    box_rows = _rows(data_root / "物资需求与配送时限.xlsx", "逐箱货箱清单")
    comm_rows = _rows(data_root / "通信链路参数.xlsx", "数据")
    return {
        "center_rows": center_rows,
        "aircraft_rows": aircraft_rows,
        "relay_rows": relay_rows,
        "box_rows": box_rows,
        "comm_rows": comm_rows,
    }


def _synthetic_terrain() -> Terrain:
    """A 5x5 DEM with a high central ridge for the midpoint-obstruction test."""

    xs = [0.0, 250.0, 500.0, 750.0, 1_000.0]
    ys = [0.0, 250.0, 500.0, 750.0, 1_000.0]
    elevations = [[0.0 for _ in xs] for _ in ys]
    elevations[2][2] = 160.0
    elevations[1][2] = 140.0
    elevations[3][2] = 140.0
    return Terrain(xs_m=xs, ys_m=ys, elevations_m=elevations)


def _read_actual_dem(root: Path, lon0: float, lat0: float) -> Terrain | None:
    try:
        from scipy.io import loadmat

        dem_path = root / GEO_DIR / "数字高程模型数据（DEM）" / "镇龙乡及周边30米DEM.mat"
        values = loadmat(dem_path)
        dem = values["dem"]
        lats = values["latitude"].reshape(-1)
        lons = values["longitude"].reshape(-1)
        xs = [_local_xy(float(lon), lat0, lon0, lat0)[0] for lon in lons]
        ys = [_local_xy(lon0, float(lat), lon0, lat0)[1] for lat in lats]
        elevation_rows = dem.tolist()
        return Terrain(xs_m=xs, ys_m=ys, elevations_m=elevation_rows)
    except Exception:
        return None


def load_mini_scenario(project_root: Path) -> tuple[Scenario, Terrain | None]:
    tables = _read_base_tables(project_root)
    center_rows = tables["center_rows"]
    aircraft_rows = tables["aircraft_rows"]
    relay_rows = tables["relay_rows"]
    box_rows = tables["box_rows"]
    comm_rows = tables["comm_rows"]

    center = next(row for row in center_rows if row[0] == "O01")
    lon0, lat0, ground0 = float(center[2]), float(center[3]), float(center[4])
    nodes: dict[str, Node] = {
        "O01": Node("O01", 0.0, 0.0, ground0),
    }
    services = ["S001", "S002", "S015"]
    for row in center_rows:
        if row[0] in services:
            x, y = _local_xy(float(row[2]), float(row[3]), lon0, lat0)
            nodes[row[0]] = Node(row[0], x, y, float(row[4]))

    type_rows = {
        row[0]: row
        for row in aircraft_rows
        if row[0] in {"A", "B", "C"}
        and isinstance(row[1], str)
        and len(row) >= 18
        and row[3] is not None
        and row[8] is not None
    }
    aircraft_types = {
        type_id: AircraftType(
            type_id=type_id,
            max_mass_kg=float(row[3]),
            max_volume_m3=float(row[4]),
            energy_kwh=float(row[8]),
            reserve_ratio=float(row[9]) / 100.0,
            charge_time_s=1800.0 if type_id == "A" else 2400.0 if type_id == "B" else 3000.0,
            max_climb_mps=float(row[14]),
            max_descend_mps=float(row[15]),
            cruise_mps=float(row[5]),
        )
        for type_id, row in type_rows.items()
    }
    aircraft = {
        "U01": Aircraft("U01", "A"),
        "U05": Aircraft("U05", "B"),
        "U07": Aircraft("U07", "C"),
    }
    batteries = {
        "A-BAT-01": Battery("A-BAT-01", "A", 1800.0),
        "B-BAT-01": Battery("B-BAT-01", "B", 2400.0),
        "C-BAT-01": Battery("C-BAT-01", "C", 3000.0),
    }

    relay_row = next(row for row in relay_rows if row[0] == "R")
    relay_type = RelayType(
        type_id="R",
        energy_kwh=float(relay_row[7]),
        reserve_ratio=float(relay_row[8]) / 100.0,
        charge_time_s=1800.0,
        max_hover_height_m=float(relay_row[18]),
    )
    relays = {"R01": Relay("R01", "R")}
    relay_components = {"R-COMP-01": RelayComponent("R-COMP-01", "R", 1800.0)}

    comm: dict[str, float] = {}
    for row in comm_rows:
        if len(row) < 5 or row[4] is None:
            continue
        if not isinstance(row[4], (int, float)):
            continue
        category = str(row[0])
        symbol = str(row[3]) if row[3] else ""
        value = float(row[4])
        if symbol in {"f", "Lsys", "Lobs", "Psens", "M"}:
            comm[symbol] = value
        elif category == "运输无人机" and symbol == "Pt":
            comm["Pt_transport"] = value
        elif category == "运输无人机" and symbol == "G":
            comm["G_transport"] = value
        elif category == "中继接入端" and symbol == "Pt":
            comm["Pt_relay_access"] = value
        elif category == "中继接入端" and symbol == "G":
            comm["G_relay_access"] = value
        elif category == "中继回传端" and symbol == "Pt":
            comm["Pt_relay_backhaul"] = value
        elif category == "中继回传端" and symbol == "G":
            comm["G_relay_backhaul"] = value
        elif category == "固定网关 G01" and symbol == "Pt":
            comm["Pt_gateway"] = value
        elif category == "固定网关 G01" and symbol == "G":
            comm["G_gateway"] = value

    selected_ids = {
        "S001-MED-01",
        "S001-WAT-01",
        "S001-FOD-01",
        "S001-HYG-01",
        "S001-HYG-02",
        "S002-MED-01",
        "S002-WAT-01",
        "S015-MED-01",
    }
    boxes: dict[str, Box] = {}
    for row in box_rows:
        if row[0] in selected_ids:
            boxes[str(row[0])] = Box(
                box_id=str(row[0]),
                service_id=str(row[1]),
                material_type=str(row[2]),
                mass_kg=float(row[3]),
                volume_m3=float(row[4]),
                first_batch=str(row[5]) == "是",
                first_deadline_s=_as_float(row[6]),
                expected_time_s=_as_float(row[7]),
                priority=float(row[8]),
            )
    if set(boxes) != selected_ids:
        missing = sorted(selected_ids - set(boxes))
        raise ValueError(f"小实例货箱缺失: {missing}")

    inventory = {"A": 6, "B": 4, "C": 4, "R": 2, "R_COMPONENT": 6}
    synthetic = _synthetic_terrain()
    actual_terrain = _read_actual_dem(project_root, lon0, lat0)
    scenario = Scenario(
        nodes=nodes,
        boxes=boxes,
        aircraft_types=aircraft_types,
        aircraft=aircraft,
        batteries=batteries,
        relay_type=relay_type,
        relays=relays,
        relay_components=relay_components,
        communication=comm,
        terrain=synthetic,
        services=services,
        inventory=inventory,
        metadata={
            "selected_box_count": len(boxes),
            "selected_services": services,
            "energy_formula_status": "BLOCKED",
            "actual_dem_loaded": actual_terrain is not None,
        },
    )
    return scenario, actual_terrain
