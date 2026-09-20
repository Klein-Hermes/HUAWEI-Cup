"""Fill only the waveform-classification cells in a copy of appendix 4."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "data" / "raw" / "2024_C题" / "附件四（Excel表）.xlsx"
PREDICTIONS = ROOT / "results" / "2024_C题_q1" / "附件二_波形分类结果.csv"
OUTPUT = ROOT / "results" / "2024_C题_q1" / "附件四（问题1已填写）.xlsx"


def main() -> None:
    prediction = pd.read_csv(PREDICTIONS)
    if len(prediction) != 80 or prediction.sample_id.tolist() != list(range(1, 81)):
        raise ValueError("Classification result does not cover sample IDs 1–80 in order")
    if not prediction.waveform_code.isin([1, 2, 3]).all():
        raise ValueError("Unexpected waveform classification code")

    book = load_workbook(TEMPLATE)
    if book.sheetnames != ["Sheet1", "Sheet2", "Sheet3"]:
        raise ValueError("Unexpected appendix-4 template sheets")
    sheet = book["Sheet1"]
    if sheet.max_row != 401 or sheet.max_column != 3 or sheet["A2"].value != 1:
        raise ValueError("Unexpected appendix-4 template layout")
    if any(sheet.cell(row, 2).value is not None for row in range(2, 402)):
        raise ValueError("Template classification column is not blank")
    for sample_id, code in zip(prediction.sample_id, prediction.waveform_code):
        sheet.cell(int(sample_id) + 1, 2).value = int(code)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    book.save(OUTPUT)

    check = load_workbook(OUTPUT)
    out = check["Sheet1"]
    actual = [out.cell(row, 2).value for row in range(2, 82)]
    expected = prediction.waveform_code.astype(int).tolist()
    if actual != expected or any(out.cell(row, 2).value is not None for row in range(82, 402)):
        raise ValueError("Saved classification cells do not match the predictions")
    if [out.cell(row, 3).value for row in range(1, 402)] != [sheet.cell(row, 3).value for row in range(1, 402)]:
        raise ValueError("The loss-prediction column changed")
    print(f"Saved {OUTPUT}; 80 classification cells filled; remaining B cells blank.")


if __name__ == "__main__":
    main()
